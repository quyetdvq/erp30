#!/usr/bin/env python3
"""Generate static Product import templates from the same versioned schema used by PHP."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "backend/app/Modules/Import/Templates/product-v1.json"
TARGET = ROOT / "frontend/public/import-templates"


def build() -> None:
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    fields = data["columns"]
    TARGET.mkdir(parents=True, exist_ok=True)
    csv_path = TARGET / "product-v1.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        csv.writer(handle, lineterminator="\r\n").writerow(fields)

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(fields)
    ws.freeze_panes = "B2"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f"A1:AE1001"
    for column, cell in enumerate(ws[1], start=1):
        cell.fill = PatternFill("solid", fgColor="184A6B")
        cell.font = Font(name="Aptos", color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = min(max(len(str(cell.value)) + 3, 17), 29)
    ws.row_dimensions[1].height = 38
    # Preserve codes, exact DECIMAL strings, and source-key IDs; never auto-format to float/date.
    for name, dim in ws.column_dimensions.items():
        dim.number_format = "@"
    validation = DataValidation(type="list", formula1='"PRODUCT,UOM_CONVERSION,COMBO_ITEM"')
    validation.promptTitle = "Record type"
    validation.prompt = "Use one canonical record type"
    validation.error = "Unsupported record_type"
    validation.errorTitle = "Invalid record_type"
    validation.showInputMessage = True
    validation.showErrorMessage = True
    ws.add_data_validation(validation)
    validation.add("A2:A1001")

    # Keep the importable Data sheet empty. Examples belong on a separate
    # worksheet, so downloading a template never imports demonstration records.
    examples = wb.create_sheet("Examples")
    examples.sheet_view.showGridLines = False
    examples.append(["DEMONSTRATION ONLY - copy validated rows into Data, do not import this sheet"])
    examples.append(fields)
    examples_data = [
        {"record_type":"PRODUCT","source_key":"sample-product","product_code":"DEMO-001",
         "sku":"DEMO-SKU-001","product_name":"Example product","product_type":"STANDARD",
         "base_uom_code":"PCS","price":"120.00"},
        {"record_type":"PRODUCT","source_key":"sample-combo","product_code":"DEMO-COMBO",
         "sku":"DEMO-SKU-BUNDLE","product_name":"Example bundle","product_type":"COMBO",
         "base_uom_code":"PCS","price":"500.00"},
        {"record_type":"UOM_CONVERSION","product_source_key":"sample-product","uom_code":"BOX",
         "conversion_factor":"12.000000","usage_type":"SALES","price":"1440.00"},
        {"record_type":"COMBO_ITEM","combo_source_key":"sample-combo",
         "component_source_key":"sample-product","quantity":"2.000000"},
    ]
    for record in examples_data:
        examples.append([record.get(field, "") for field in fields])
    examples.freeze_panes = "B3"
    for cell in examples[2]:
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="184A6B")
    for cell in examples[1]:
        cell.font = Font(name="Aptos", bold=True, color="A32020")
    for cell in examples[2]:
        examples.column_dimensions[cell.column_letter].width = min(max(len(str(cell.value)) + 3, 17), 29)
    examples.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(fields))
    examples.row_dimensions[1].height=26

    help_sheet = wb.create_sheet("Guide")
    help_sheet.sheet_view.showGridLines = False
    notes = [
        ("ERP Import Template", "PRODUCT v1 | CSV/XLSX"),
        ("Input", "Fill only Data sheet. Do not rename or duplicate headers."),
        ("Examples", "Examples worksheet is demonstration only; copied codes and UOMs may not exist on your server."),
        ("Mode", "CREATE_ONLY | UPDATE_ONLY | UPSERT (selected in UI)"),
        ("PRODUCT", "Required source_key and product_name for new root products."),
        ("UOM_CONVERSION", "Refer to an existing Product code or product_source_key."),
        ("COMBO_ITEM", "Specify COMBO owner and component by source_key or code."),
        ("Source key", "Internal file reference; it is not the persisted product_code."),
        ("Null", "Use __NULL__ for an explicit null value; empty cells mean omitted."),
        ("Decimals", "Enter decimal values as text (e.g. 1.250000), never formulas."),
        ("Boolean", "Use 1/0 or true/false as plain text for serial_management_enabled."),
        ("Lifecycle", "Status changes require the set-state API; import does not bypass it."),
        ("Permissions", "product.import and target create/edit PBAC are both checked."),
        ("Limit", f"Up to {data['max_rows']} rows, {data['max_bytes']//1024} KiB per file."),
        ("Version", "1 (change the template version only with a reviewed adapter update)"),
    ]
    for label, note in notes:
        help_sheet.append([label, note])
    for cell in help_sheet[1]:
        cell.fill = PatternFill("solid", fgColor="184A6B")
        cell.font = Font(color="FFFFFF", bold=True)
    help_sheet.column_dimensions["A"].width = 23
    help_sheet.column_dimensions["B"].width = 93
    for row in help_sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    wb.save(TARGET / "product-v1.xlsx")

    with csv_path.open(newline="", encoding="utf-8") as handle:
        assert next(csv.reader(handle)) == fields
    verify = load_workbook(TARGET / "product-v1.xlsx", read_only=True, data_only=False)
    assert [cell.value for cell in next(verify["Data"].rows)] == fields
    assert verify.sheetnames == ["Data", "Examples", "Guide"]
    assert verify["Data"].max_row == 1, "Never ship populated import Data rows"
    verify.close()
    print(f"PASS: {len(fields)} canonical columns in CSV and XLSX")


if __name__ == "__main__":
    build()
